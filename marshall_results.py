"""Cálculos del VBA de rota 1.xls / rota 2.xls, sin ejecutar macros."""
import math
from une_marshall import LABELS,STANDARD,STANDARD_TITLE,STANDARD_URL,normative

LEGACY_LABELS=[('peak_kp','Estabilidad máxima (kp)'),('load95_kp','95 % de estabilidad (kp)'),
 ('def50_mm','Deformación a 50 kp (mm)'),('def95_mm','Deformación al 95 % (mm)'),
 ('marshall_mm','Deformación Marshall: d95 − d50 (mm)'),('slope_kp_mm','Pendiente tangente (kp/mm)'),
 ('zero_mm','Cero de la tangente (mm)'),('peak_mm','Deformación al máximo (mm)'),
 ('corrected_mm','Deformación corregida: dmáx − cero (mm)')]

def calculate(rows):
    rows=[dict(r) for r in rows if math.isfinite(r['force_kp']) and math.isfinite(r['deformation_mm'])]
    result={key:None for key,_ in LEGACY_LABELS};result.update(rows=rows,tangent=[],differences=[],mean_differences=[])
    if not rows:return normative(result)
    x=[r['deformation_mm'] for r in rows];y=[r['force_kp'] for r in rows]
    peak=max(range(len(y)),key=y.__getitem__)
    if y[peak]<=0:return normative(result)
    result.update(peak_kp=y[peak],load95_kp=math.floor(.95*y[peak]),peak_mm=x[peak])
    truncate=lambda v:math.floor(v*100+1e-9)/100
    def crossing(target,last=False):
        pairs=[i for i in range(1,peak+1) if y[i-1]<target<=y[i] and y[i]>y[i-1]]
        if not pairs:return None
        i=pairs[-1] if last else pairs[0]
        return truncate(x[i-1]+(x[i]-x[i-1])*(target-y[i-1])/(y[i]-y[i-1]))
    result['def50_mm']=crossing(50,last=True);result['def95_mm']=crossing(result['load95_kp'])
    if result['def50_mm'] is not None and result['def95_mm'] is not None:
        result['marshall_mm']=round(result['def95_mm']-result['def50_mm'],2)
    differences=[None]+[b-a for a,b in zip(y,y[1:])]
    means=[None]*len(rows)
    for i in range(3,len(rows)-2):means[i]=sum(differences[i-2:i+3])/5
    result.update(differences=differences,mean_differences=means)
    candidates=[i for i in range(3,min(peak+1,len(rows)-2)) if means[i] is not None and means[i]>0 and x[i+2]>x[i-2]]
    if candidates:
        i=max(candidates,key=lambda j:means[j])
        slope=(y[i+2]-y[i-2])/(x[i+2]-x[i-2])
        if slope>0:
            zero=x[i]-y[i]/slope
            result.update(slope_kp_mm=slope,zero_mm=zero,corrected_mm=truncate(x[peak]-zero),
                tangent=[(zero,0),(zero+y[peak]/slope,y[peak])],tangent_index=i)
    return normative(result)


def calculation_rows(session):
    """Exactamente los puntos usados por pantalla y Excel, bloques de 10."""
    from plot_cleaning import compression_deformation
    raw=session['samples']
    if (session.get('configuration') or {}).get('export_processed'):
        counts=session['configuration'].get('export_block_counts',[])
        return [{**r,'samples_count':counts[i] if i<len(counts) else r.get('samples_count',1)} for i,r in enumerate(raw)]
    imported=all(r.get('vf') is None or not math.isfinite(r['vf']) for r in raw)
    rows=[]
    if imported:
        rows=[{**r,'samples_count':1} for r in raw if r['force_kp']>=0]
    else:
        for start in range(0,len(raw),10):
            block=raw[start:start+10]
            accepted=[r['force_kp'] for r in block if math.isfinite(r['force_kp']) and r['force_kp']>=0]
            row=dict(t=block[-1]['t'],force_kp=sum(accepted)/len(accepted) if accepted else 0.,
                deformation_mm=sum(r['deformation_mm'] for r in block)/len(block),samples_count=len(block))
            for key in ('vf','vd'):
                values=[r[key] for r in block if r.get(key) is not None and math.isfinite(r[key]) and (key!='vf' or r['force_kp']>=0)]
                row[key]=sum(values)/len(values) if values else None
            rows.append(row)
    deformation=compression_deformation([r['deformation_mm'] for r in rows])
    return [{**r,'deformation_mm':d} for r,d in zip(rows,deformation)]


def session_results(session):
    result=calculate(calculation_rows(session))
    result['processing']='Media de 10 lecturas por punto; último bloque parcial incluido' if (session.get('configuration') or {}).get('export_processed') or any(r.get('vf') is not None for r in session['samples']) else 'TXT ya procesado; sin segundo promedio'
    cfg=session.get('configuration') or {}
    return normative(result,cfg.get('specimen_height_mm'),cfg.get('specimen_volume_mm3'))


def draw_results(ax,result):
    ax.clear(); ax.set_title('Fuerza–deformación y tangente Marshall');ax.set_xlabel('Deformación (mm)');ax.set_ylabel('Fuerza (kp)');ax.grid(alpha=.2)
    rows=result['rows']
    if rows:ax.plot([r['deformation_mm'] for r in rows],[r['force_kp'] for r in rows],label='Ensayo',color='#14616b')
    if rows:
        peak=max(range(len(rows)),key=lambda i:rows[i]['force_kp'])
        ax.plot([r['deformation_mm'] for r in rows[peak:]],[r['force_kp'] for r in rows[peak:]],color='#c62828',label='Desde carga máxima')
    if result['tangent']:
        x,y=zip(*result['tangent']);ax.plot(x,y,'--',color='#d58921',label='Tangente de máxima subida')
    if result.get('zero_mm') is not None:ax.axvline(result['zero_mm'],color='#888888',linestyle=':',label='A: cero de tangente')
    if rows:
        ax.legend(fontsize=8)
        peak=max(rows,key=lambda r:r['force_kp'])
        ax.scatter([peak['deformation_mm']],[peak['force_kp']],color='#c62828',s=25,zorder=5)
    ax.set_ylim(bottom=0)


def add_excel_results(workbook,session):
    """Curva XY y valores UNE; h/v editables con fórmulas nativas de Excel."""
    from openpyxl.chart import ScatterChart,Series,Reference
    from openpyxl.chart.axis import ChartLines
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.properties import CalcProperties
    result=session_results(session);cfg=session.get('configuration') or {}
    sheet=workbook.create_sheet('Resultados Marshall',0)
    graph=workbook.create_sheet('Gráfico Marshall',1)
    reference=workbook.create_sheet('Referencia UNE',2)
    headers=['Deformación (mm)','Carga media (kp)','Tiempo (s)','Diferencia carga (kp)','Media 5 diferencias (kp)','Lecturas por punto']
    sheet.append(headers)
    for i,r in enumerate(result['rows']):
        sheet.append([r['deformation_mm'],r['force_kp'],r['t'],result['differences'][i] if result['differences'] else None,result['mean_differences'][i] if result['mean_differences'] else None,r.get('samples_count',1)])
    sheet['G1']='Resultado UNE';sheet['H1']='Valor'
    for i,(key,label) in enumerate(LABELS,2):
        sheet.cell(i,7,label);sheet.cell(i,8,result.get(key) if result.get(key) is not None else 'Pendiente h/v' if key in ('stability_kn','correction_factor','quotient_kn_mm') else 'No disponible')
    sheet['G12']='Referencia';sheet['H12']=STANDARD
    sheet['G13']='Apartados';sheet['H13']=result['standard_sections']
    sheet['G14']='Procesado';sheet['H14']=result['processing']
    sheet['G16']='Altura h (mm)';sheet['H16']='=IF(ISNUMBER(\'Gráfico Marshall\'!R20),\'Gráfico Marshall\'!R20,"")'
    sheet['G17']='Volumen v (mm³)';sheet['H17']='=IF(ISNUMBER(\'Gráfico Marshall\'!R21),\'Gráfico Marshall\'!R21,"")'
    sheet['G18']='Factor c, 7.1 fórmula (1)'
    sheet['H18']='=IF(AND(ISNUMBER(H16),H16>0),5.2*EXP(-0.0259*H16),IF(AND(ISNUMBER(H17),H17>0),5.2*EXP(-0.0000032*H17),""))'
    sheet['G20']='Carga máxima exacta (kN)';sheet['H20']=result.get('raw_stability_exact_kn')
    sheet['G21']='F exacta (mm)';sheet['H21']=result.get('flow_exact_mm')
    sheet['G22']='Pendiente automática (kp/mm)';sheet['H22']=result.get('slope_kp_mm')
    sheet['H4']='=IF(AND(ISNUMBER(H18),ISNUMBER(H20)),ROUND(H20*H18,1),"Pendiente h/v")'
    sheet['H8']='=IF(AND(ISNUMBER(H18),ISNUMBER(H20),ISNUMBER(H21),H21>0),ROUND(H20*H18/H21,1),"Pendiente h/v o F")'
    sheet['H10']='=IF(ISNUMBER(H18),H18,"Pendiente h/v")'
    sheet['G24']='Tangente X (mm)';sheet['H24']='Tangente Y (kp)'
    for i,(x,y) in enumerate(result['tangent'],25):sheet.cell(i,7,x);sheet.cell(i,8,y)
    graph.merge_cells('A1:S1');graph['A1']='Resultados Marshall'
    graph.merge_cells('A2:S2');graph['A2']='Ensayo: '+session['metadata'].get('reference','');graph['A2'].data_type='s'
    graph.merge_cells('A3:S3');graph['A3']=STANDARD+' · '+result['standard_sections'];graph['A3'].hyperlink=STANDARD_URL
    graph.merge_cells('A4:S4');graph['A4']=''
    for i,(key,label) in enumerate(LABELS,6):
        graph.merge_cells(start_row=i,start_column=13,end_row=i,end_column=17)
        graph.cell(i,13,label)
        graph.merge_cells(start_row=i,start_column=18,end_row=i,end_column=19)
        source=f"'Resultados Marshall'!H{i-4}"
        graph.cell(i,18,f'=IF({source}="","No disponible",{source})')
        graph.cell(i,18).number_format='0' if key=='peak_kp' else '0.0000' if key=='correction_factor' else '0.00' if key=='zero_mm' else '0.0'
        graph.cell(i,18).font=Font(size=12,bold=True,color='14505B')
        graph.cell(i,18).alignment=Alignment(horizontal='right',vertical='center')
        graph.cell(i,13).alignment=Alignment(vertical='center',wrap_text=True)
        graph.row_dimensions[i].height=29
        for col in range(13,20):graph.cell(i,col).fill=PatternFill('solid',fgColor='F1F6F6' if i%2==0 else 'FFFFFF')
    graph.merge_cells('M18:S18');graph['M18']='Datos de probeta · rellenar una opción'
    for row,label,key in ((20,'Altura h (mm)','specimen_height_mm'),(21,'Volumen v (mm³)','specimen_volume_mm3')):
        graph.merge_cells(start_row=row,start_column=13,end_row=row,end_column=17);graph.cell(row,13,label)
        graph.merge_cells(start_row=row,start_column=18,end_row=row,end_column=19)
        cell=graph.cell(row,18,cfg.get(key));cell.fill=PatternFill('solid',fgColor='FFF2CC');cell.number_format='0.0' if row==20 else '0'
        cell.border=Border(bottom=Side(style='thin',color='C59B30'))
        validation=DataValidation(type='decimal',operator='greaterThan',formula1=0,allow_blank=True)
        validation.errorTitle='Dimensión inválida';validation.error='Introduzca un número positivo.';validation.showErrorMessage=True
        graph.add_data_validation(validation);validation.add(cell)
    graph.merge_cells('M23:S23');graph['M23']='Si rellena h y v, se usa h.'
    graph.merge_cells('M24:S24');graph['M24']='Introduzca una dimensión para calcular S y S/F.'
    graph.merge_cells('M26:S28');graph['M26']='=IF(AND(ISNUMBER(R20),R20>0,OR(R20<60.5,R20>66.5)),"Altura fuera de 60,5–66,5 mm: revisar precisión (§7.1).","")';graph['M26'].alignment=Alignment(wrap_text=True)
    graph.merge_cells('M30:S30');graph['M30']='Resultados individuales (§7.3 y §8).'
    graph.merge_cells('M31:S31');graph['M31']='Revisar el grupo y las condiciones del ensayo.'
    chart=ScatterChart();chart.title='Fuerza–deformación';chart.x_axis.title='Deformación (mm)';chart.y_axis.title='Fuerza (kp)'
    chart.width=21;chart.height=13;chart.scatterStyle='line';chart.display_blanks='gap';chart.legend.position='b'
    n=len(result['rows'])
    if n:
        chart.series.append(Series(Reference(sheet,min_col=2,min_row=2,max_row=n+1),Reference(sheet,min_col=1,min_row=2,max_row=n+1),title='Ensayo'))
        chart.series[0].graphicalProperties.line.solidFill='14616B'
        peak=max(range(n),key=lambda i:result['rows'][i]['force_kp'])
        descent=Series(Reference(sheet,min_col=2,min_row=peak+2,max_row=n+1),Reference(sheet,min_col=1,min_row=peak+2,max_row=n+1),title='Desde carga máxima')
        descent.graphicalProperties.line.solidFill='C62828';chart.series.append(descent)
    if result['tangent']:
        tangent=Series(Reference(sheet,min_col=8,min_row=25,max_row=26),Reference(sheet,min_col=7,min_row=25,max_row=26),title='Tangente estimada')
        tangent.graphicalProperties.line.solidFill='D58921';tangent.graphicalProperties.line.prstDash='dash';chart.series.append(tangent)
    for series in chart.series:series.marker.symbol='none'
    for axis,position,fmt in ((chart.x_axis,'b','0.0'),(chart.y_axis,'l','0')):
        axis.axPos=position;axis.scaling.min=0;axis.tickLblPos='nextTo';axis.majorTickMark='out';axis.delete=False;axis.numFmt=fmt
    chart.y_axis.majorGridlines=ChartLines();graph.add_chart(chart,'A6')
    notes=[['Referencia normativa',STANDARD],['Título',STANDARD_TITLE],['Edición','Marzo de 2022; versión española de EN 12697-34:2020'],['Fuente',STANDARD_URL],
        ['Fuente de cálculo','PDF aportado por el usuario, apartados 3, 7 y anexo A.'],
        ['S corregida (kN)','S = carga máxima (kN) × c'],['c por altura h (mm)','c = 5,2 × exp(−0,0259 × h)'],['c por volumen v (mm³)','c = 5,2 × exp(−3,2 × 10^−6 × v)'],
        ['F (mm)','F = deformación al máximo − corte A de la tangente'],['Ft (mm)','Ft = carga máxima / pendiente de la tangente'],['FT (mm)','FT = deformación total al máximo'],['Cociente Marshall (kN/mm)','S / F; se calcula antes del redondeo de presentación'],
        ['Redondeo','S: 0,1 kN; F, Ft, FT: 0,1 mm; S/F: 0,1 kN/mm (§7.1–7.4).'],
        ['Tangente automática','Estimación numérica local en la rama ascendente. Revisar sobre la curva; la norma no prescribe este algoritmo.'],
        ['Alcance','Datos de una probeta. No certifica por sí solo calibración, velocidad, temperatura o validez de un grupo de cuatro probetas.'],
        ['Procesado',result['processing']],['Número de puntos dibujados y exportados',n],
        ['Complementario histórico, fuera de UNE','Los valores d50/d95 del VBA antiguo se conservan solo como información auxiliar.'],
        ['d50 (mm)',result.get('def50_mm')],['d95 (mm)',result.get('def95_mm')],['d95 − d50 (mm)',result.get('marshall_mm')]]
    for row in notes:reference.append(row)
    reference['B4'].hyperlink=STANDARD_URL
    for ws in (sheet,reference):
        ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='14505B')
    for col in ('A','B','C','D','E','F'):sheet.column_dimensions[col].width=23
    sheet.column_dimensions['G'].width=46;sheet.column_dimensions['H'].width=48
    for row in sheet.iter_rows(min_row=2,max_col=8):
        for cell in row:
            if isinstance(cell.value,(int,float)):cell.number_format='0.0000' if cell.column==8 else '0.00'
    for row in (3,4,5,6,7,8):sheet.cell(row,8).number_format='0.0'
    sheet['H2'].number_format='0';sheet['H10'].number_format='0.0000'
    for row in (12,13,14):sheet.cell(row,8).alignment=Alignment(wrap_text=True);sheet.row_dimensions[row].height=32
    sheet.row_dimensions[1].height=34
    for cell in sheet[1]:cell.alignment=Alignment(wrap_text=True,vertical='center')
    for col in 'ABCDEFGHIJKLMNOPQRS':graph.column_dimensions[col].width=10.5
    graph.column_dimensions['L'].width=3
    graph['A1'].font=Font(size=17,bold=True,color='14505B');graph['A2'].font=Font(size=11,color='55636B');graph['A3'].font=Font(size=10,color='14505B');graph['A4'].font=Font(size=10,italic=True,color='55636B')
    graph.row_dimensions[1].height=28;graph.row_dimensions[3].height=20;graph['M18'].font=Font(size=11,bold=True,color='14505B')
    reference.column_dimensions['A'].width=48;reference.column_dimensions['B'].width=100
    for row in reference.iter_rows(min_row=2):
        row[1].alignment=Alignment(wrap_text=True,vertical='top');reference.row_dimensions[row[0].row].height=32
    graph.sheet_view.showGridLines=False;graph.sheet_properties.pageSetUpPr.fitToPage=True
    graph.page_setup.orientation='landscape';graph.page_setup.paperSize=graph.PAPERSIZE_A4;graph.page_setup.fitToWidth=1;graph.page_setup.fitToHeight=1;graph.print_area='A1:S37'
    workbook.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
    return result
