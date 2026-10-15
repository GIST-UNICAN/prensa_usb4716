import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from app import Sample
from sessions import make_session, SessionStore, load_session, validate_session, session_summary
from exports import export_session, chart_series, moving_average

class SessionTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.root=Path(self.tmp.name)
        self.session=make_session([Sample(0,0,0,0,0),Sample(.5,-1,.5,100,1),Sample(1,-2,1,200,2)],
                                  dict(reference='Probeta 1',sample_id='A-1',operator='Ana',notes='Control',demo=True))
    def tearDown(self): self.tmp.cleanup()

    def test_history_survives_reopen_and_updates_same_id(self):
        db=self.root/'history.sqlite3'; store=SessionStore(db); store.save(self.session)
        self.session['metadata']['notes']='Actualizado'; store.save(self.session)
        reopened=SessionStore(db)
        self.assertEqual(len(reopened.list()),1)
        self.assertEqual(reopened.get(self.session['id'])['metadata']['notes'],'Actualizado')
        self.assertEqual(len(reopened.list('ana','Simulación')),1)
        self.assertEqual(reopened.list('Ana','Real'),[])
        self.assertEqual(session_summary(self.session)['deformation_at_peak_mm'],2)

    def test_delete_is_persistent_and_scoped_to_selected_id(self):
        store=SessionStore(self.root/'db');store.save(self.session)
        other=json.loads(json.dumps(self.session));other['id']='other';store.save(other)
        store.delete(self.session['id'])
        self.assertEqual([s['id'] for s in SessionStore(self.root/'db').list()],['other'])
        with self.assertRaises(ValueError):store.delete(self.session['id'])

    def test_invalid_session_cannot_replace_valid_archive(self):
        store=SessionStore(self.root/'db'); store.save(self.session)
        self.session['samples'][1]['force_kp']=float('nan')
        with self.assertRaises(ValueError): store.save(self.session)
        self.assertEqual(store.get(self.session['id'])['samples'][1]['force_kp'],100)

    def test_unknown_voltages_json_roundtrip(self):
        session=make_session([Sample(0,float('nan'),float('nan'),5,1)],{'reference':'Antiguo'})
        path=self.root/'session.json'; export_session(session,path,'json')
        loaded=load_session(path)
        self.assertIsNone(loaded['samples'][0]['vf'])
        self.assertNotIn('NaN',path.read_text())
        series,_,_=chart_series(loaded,'volts')
        self.assertEqual(series[0][1],[None])

    def test_bad_order_and_version_rejected(self):
        self.session['samples'][2]['t']=.1
        with self.assertRaises(ValueError): validate_session(self.session)
        self.session['version']=999
        with self.assertRaises(ValueError): validate_session(self.session)

    def test_speed_and_smoothing_leave_samples_unchanged(self):
        before=json.dumps(self.session)
        series,_,_=chart_series(self.session,'speed')
        self.assertEqual(series[0][1],[120,120])
        self.assertEqual(moving_average([0,10,20],2),[0,5,15])
        self.assertEqual(json.dumps(self.session),before)
        series,_,_=chart_series(self.session,'ft','kN')
        self.assertAlmostEqual(series[0][1][-1],1.96133)

    def test_xlsx_contains_data_and_charts_and_literal_metadata(self):
        from openpyxl import load_workbook
        self.session['metadata']['reference']='=1+1'
        path=self.root/'book.xlsx'; export_session(self.session,path,'xlsx')
        book=load_workbook(path)
        self.assertIn('Resultados Marshall',book.sheetnames)
        self.assertIn('Gráfico Marshall',book.sheetnames)
        self.assertIn('Referencia UNE',book.sheetnames)
        self.assertEqual(book['Referencia UNE']['B1'].value,'UNE-EN 12697-34:2022')
        self.assertIsNone(book['Gráfico Marshall']['R20'].value)
        self.assertIn('EXP(-0.0259',book['Resultados Marshall']['H18'].value)
        chart=book['Gráfico Marshall']._charts[0]
        self.assertEqual(chart.x_axis.tickLblPos,'nextTo');self.assertEqual(chart.y_axis.tickLblPos,'nextTo')
        self.assertFalse(chart.x_axis.delete);self.assertFalse(chart.y_axis.delete)
        self.assertIsInstance(book['Resultados Marshall']['H2'].value,(int,float))
        self.assertEqual(len(book['Gráfico Marshall']._charts),1)
        self.assertEqual(book['Resultados Marshall']['F2'].value,3)
        self.assertIsNone(book['Resultados Marshall']['A3'].value)
        self.assertNotIn('Datos',book.sheetnames)
        self.assertEqual(book['Configuración']['B2'].data_type,'s')
        book.close()

    def test_xlsx_prefills_optional_geometry(self):
        from openpyxl import load_workbook
        self.session['configuration']={'specimen_height_mm':63.5,'specimen_volume_mm3':500000}
        path=self.root/'dimensions.xlsx';export_session(self.session,path,'xlsx')
        book=load_workbook(path)
        self.assertEqual(book['Gráfico Marshall']['R20'].value,63.5)
        self.assertEqual(book['Gráfico Marshall']['R21'].value,500000)
        book.close()

    def test_zip_complete_and_html_escapes_metadata(self):
        self.session['metadata']['notes']='<script>alert(1)</script>'
        path=self.root/'package.zip'; export_session(self.session,path,'zip')
        with zipfile.ZipFile(path) as archive:
            self.assertEqual(len([x for x in archive.namelist() if x.endswith('.png')]),5)
            self.assertEqual(validate_session(json.loads(archive.read('ensayo.json')))['samples'],self.session['samples'])
            report=archive.read('informe.html').decode()
            self.assertIn('&lt;script&gt;',report)
            self.assertNotIn('<script>',report)
            self.assertEqual(report.count('<svg'),5)
            self.assertIn('SIMULACIÓN',report)

    def test_failed_export_preserves_destination(self):
        path=self.root/'report.html'; path.write_text('original')
        with patch('exports.html_report',side_effect=RuntimeError('disk/renderer')):
            with self.assertRaises(RuntimeError): export_session(self.session,path,'html')
        self.assertEqual(path.read_text(),'original')
        self.assertEqual(list(self.root.glob('.marshall-*')),[])

if __name__=='__main__':unittest.main()
